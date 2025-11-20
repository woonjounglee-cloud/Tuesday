"""샘플 데이터베이스 파일 생성 스크립트"""
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    # Excel 파일 생성
    wb = Workbook()
    wb.remove(wb.active)  # 기본 시트 제거

    # ROI 시트
    ws_roi = wb.create_sheet('ROI')
    roi_headers = ['종목코드', '종목명', '초기투자금', '수량', '종가', '평가금', '수익률']
    roi_sample_data = [
        ['449170', 'TIGER KOFR금리액티브(합성)', 5504005, 52, 110212, 5731024, 4.1],
        ['411060', 'ACE KRX금현물', 1942055, 85, 27185, 2310725, 19.0],
        ['468630', 'KODEX iShares미국투자등급회사채액티브', 736740, 66, 11985, 791010, 7.4],
        ['476760', 'ACE 미국30년국채액티브', 1291165, 127, 10265, 1303655, 1.0],
        ['0046A0', 'TIGER 미국초단기(3개월이하)국채', 331495, 33, 10180, 335940, 1.3],
        ['273130', 'KODEX 종합채권(AA-이상)액티브', 446440, 4, 116270, 465080, 4.2],
        ['379800', 'KODEX 미국S&P500', 1627410, 78, 22055, 1720290, 5.7],
        ['379810', 'KODEX 미국나스닥100', 1716035, 90, 23870, 2148300, 25.2],
        ['489250', 'KODEX 미국배당다우존스', 1990000, 165, 10842, 1788930, -10.1],
        ['294400', 'KIWOOM 200TR', 443243, 10, 72105, 721050, 62.7],
        ['453870', 'TIGER 인도니프티50', 912282, 73, 14310, 1044630, 14.5]
    ]

    # ROI 헤더
    for col_idx, header in enumerate(roi_headers, 1):
        cell = ws_roi.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')

    # ROI 데이터
    for row_idx, row_data in enumerate(roi_sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws_roi.cell(row=row_idx, column=col_idx, value=value)

    # Balance 시트
    ws_balance = wb.create_sheet('Balance')
    balance_headers = ['연도', '월', '투자원금', '추가납입', '잔고', '수익률', '기타']
    balance_sample_data = [
        ['2025년', '11월', 16940870, 0, 18360634, '8.4%', '']
    ]

    # Balance 헤더
    for col_idx, header in enumerate(balance_headers, 1):
        cell = ws_balance.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')

    # Balance 데이터
    for row_idx, row_data in enumerate(balance_sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws_balance.cell(row=row_idx, column=col_idx, value=value)

    # Portfolio 시트
    ws_portfolio = wb.create_sheet('Portfolio')
    portfolio_headers = ['종목코드', '종목명', '자산군', '구분', '세팅비중', '현재비중', '빨강매수']
    portfolio_sample_data = [
        ['449170', 'TIGER KOFR금리액티브(합성)', '안전자산', '현금', 10.0, 31.2, -3894961],
        ['411060', 'ACE KRX금현물', '대체자산', '금', 15.0, 12.6, 443370],
        ['468630', 'KODEX iShares미국투자등급회사채액티브', '채권', '미국회사채', 5.0, 4.3, 127022],
        ['476760', 'ACE 미국30년국채액티브', '채권', '미국장기채', 5.0, 7.1, -385623],
        ['0046A0', 'TIGER 미국초단기(3개월이하)국채', '채권', '미국단기채', 2.5, 1.8, 123076],
        ['273130', 'KODEX 종합채권(AA-이상)액티브', '채권', '국채', 2.5, 2.5, -6064],
        ['379800', 'KODEX 미국S&P500', '주식', '미국지수', 20.0, 9.4, 1951837],
        ['379810', 'KODEX 미국나스닥100', '주식', '미국지수', 20.0, 11.7, 1523827],
        ['489250', 'KODEX 미국배당다우존스', '주식', '미국배당', 10.0, 9.7, 47133],
        ['294400', 'KIWOOM 200TR', '주식', '국내지수', 5.0, 3.9, 196982],
        ['453870', 'TIGER 인도니프티50', '주식', '신흥국', 5.0, 5.7, -126598]
    ]

    # Portfolio 헤더
    for col_idx, header in enumerate(portfolio_headers, 1):
        cell = ws_portfolio.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')

    # Portfolio 데이터
    for row_idx, row_data in enumerate(portfolio_sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws_portfolio.cell(row=row_idx, column=col_idx, value=value)

    # 파일 저장
    wb.save('db/Home_IRP.xlsx')
    print('✅ Sample database file created: db/Home_IRP.xlsx')

except ImportError:
    print('⚠️  openpyxl is not installed. Please install it first.')
    print('Run: pip install openpyxl')
